"""
Génération du rapport de conformité du monitoring radio (Kiyanza - cahier
des charges IA, section 4.7).

Le cahier des charges décrit précisément ce que doit contenir ce rapport :
"pour chaque diffusion prévue, si le spot a été diffusé, à quelle heure
réellement, et l'écart éventuel avec l'heure prévue" — avec l'exemple :
"Un spot prévu à 8h05... détecté à 8h12 ; le rapport signale un décalage
de 7 minutes."

Deux formats supportés (le cahier des charges accepte les deux) : PDF
(lecture rapide, présentation à un client) et Excel (données brutes,
réutilisables pour une analyse plus poussée).

Règle de conformité : un écart de quelques minutes est normal en radio
(la diffusion réelle dépend de la régie), on ne signale pas ça comme un
problème. Seul un écart significatif ou une absence totale de diffusion
doit attirer l'attention (§4.7 : "sans fausse alerte excessive").
"""

from datetime import datetime
from pathlib import Path

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from radio_schedule import get_entries_for_report

# Écart (en minutes) en dessous duquel une diffusion est jugée "Conforme"
# malgré un léger décalage — évite de signaler comme anomalie un simple
# délai technique normal de régie radio.
COMPLIANCE_TOLERANCE_MINUTES = 5

REPORTS_DIR = Path(__file__).resolve().parent / "reports"


def _compute_row_status(entry: dict) -> str:
    if entry.get("error_message"):
        return "Erreur de capture"
    if not entry["detected"]:
        return "Non diffusé"
    deviation = float(entry["deviation_minutes"]) if entry["deviation_minutes"] is not None else 0
    if abs(deviation) <= COMPLIANCE_TOLERANCE_MINUTES:
        return "Conforme"
    return "Retard" if deviation > 0 else "Avance"


def _format_entries_for_report(entries: list) -> list:
    """Transforme les lignes brutes de la base en lignes prêtes à afficher
    (dates formatées en français, écart en texte lisible, statut calculé)."""
    rows = []
    for e in entries:
        planned = e["planned_datetime"]
        detected_at = e["detected_at"]
        deviation = e["deviation_minutes"]

        rows.append({
            "station_name": e["station_name"],
            "planned_str": planned.strftime("%d/%m/%Y %Hh%M"),
            "detected_str": detected_at.strftime("%d/%m/%Y %Hh%M") if detected_at else "—",
            "deviation_str": f"{'+' if deviation and deviation > 0 else ''}{deviation} min" if deviation is not None else "—",
            "deviation_value": float(deviation) if deviation is not None else None,
            "status": _compute_row_status(e),
        })
    return rows


def _compute_summary(rows: list) -> dict:
    total = len(rows)
    diffused = sum(1 for r in rows if r["status"] in ("Conforme", "Retard", "Avance"))
    conforme = sum(1 for r in rows if r["status"] == "Conforme")
    deviations = [r["deviation_value"] for r in rows if r["deviation_value"] is not None]

    return {
        "total": total,
        "diffused": diffused,
        "conforme": conforme,
        "taux_diffusion_pct": round(diffused / total * 100, 1) if total else 0.0,
        "taux_conformite_pct": round(conforme / total * 100, 1) if total else 0.0,
        "ecart_moyen_min": round(sum(deviations) / len(deviations), 2) if deviations else None,
    }


def generate_pdf_report(campaign_id: str = None, station_name: str = None, output_path: str = None) -> str:
    entries = get_entries_for_report(campaign_id=campaign_id, station_name=station_name)
    if not entries:
        raise ValueError("Aucune diffusion traitée trouvée pour ces critères.")

    rows = _format_entries_for_report(entries)
    summary = _compute_summary(rows)

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    if output_path is None:
        output_path = REPORTS_DIR / f"rapport_conformite_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.pdf"
    output_path = Path(output_path)

    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("Rapport de conformité — Monitoring radio", styles["Title"]))
    story.append(Paragraph(f"Généré le {datetime.utcnow().strftime('%d/%m/%Y à %Hh%M')} UTC", styles["Normal"]))
    if station_name:
        story.append(Paragraph(f"Station : {station_name}", styles["Normal"]))
    story.append(Spacer(1, 0.5 * cm))

    summary_text = (
        f"<b>{summary['diffused']} / {summary['total']}</b> diffusions détectées "
        f"({summary['taux_diffusion_pct']}%) — "
        f"<b>{summary['conforme']}</b> conformes à l'horaire prévu ({summary['taux_conformite_pct']}%)."
    )
    if summary["ecart_moyen_min"] is not None:
        summary_text += f" Écart moyen : {summary['ecart_moyen_min']:+.1f} min."
    story.append(Paragraph(summary_text, styles["Normal"]))
    story.append(Spacer(1, 0.7 * cm))

    table_data = [["Station", "Heure prévue", "Heure réelle", "Écart", "Statut"]]
    for r in rows:
        table_data.append([r["station_name"], r["planned_str"], r["detected_str"], r["deviation_str"], r["status"]])

    table = Table(table_data, repeatRows=1)
    status_colors = {
        "Conforme": colors.HexColor("#2e7d32"),
        "Retard": colors.HexColor("#c62828"),
        "Avance": colors.HexColor("#ef6c00"),
        "Non diffusé": colors.HexColor("#b71c1c"),
        "Erreur de capture": colors.HexColor("#616161"),
    }
    style_commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565c0")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
    ]
    for i, r in enumerate(rows, start=1):
        color = status_colors.get(r["status"])
        if color:
            style_commands.append(("TEXTCOLOR", (4, i), (4, i), color))
            style_commands.append(("FONTNAME", (4, i), (4, i), "Helvetica-Bold"))
    table.setStyle(TableStyle(style_commands))
    story.append(table)

    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    doc.build(story)

    return str(output_path)


def generate_excel_report(campaign_id: str = None, station_name: str = None, output_path: str = None) -> str:
    entries = get_entries_for_report(campaign_id=campaign_id, station_name=station_name)
    if not entries:
        raise ValueError("Aucune diffusion traitée trouvée pour ces critères.")

    rows = _format_entries_for_report(entries)

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    if output_path is None:
        output_path = REPORTS_DIR / f"rapport_conformite_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    output_path = Path(output_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "Conformité"

    headers = ["Station", "Heure prévue", "Heure réelle", "Écart (min)", "Statut"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        ws.append([
            r["station_name"], r["planned_str"], r["detected_str"],
            r["deviation_value"], r["status"],
        ])

    n_rows = len(rows)
    last_row = n_rows + 1

    for row_num in range(2, last_row + 1):
        for col_num in range(1, len(headers) + 1):
            ws.cell(row=row_num, column=col_num).font = Font(name="Arial")

    # Résumé, avec formules Excel (pas de valeurs codées en dur) pour que
    # le fichier se recalcule si les données sont modifiées.
    summary_start_row = last_row + 3
    status_range = f"E2:E{last_row}"
    deviation_range = f"D2:D{last_row}"

    summary_rows = [
        ("Total diffusions prévues", f"=COUNTA(A2:A{last_row})"),
        ("Diffusions détectées", f'=COUNTIF({status_range},"Conforme")+COUNTIF({status_range},"Retard")+COUNTIF({status_range},"Avance")'),
        ("Diffusions conformes", f'=COUNTIF({status_range},"Conforme")'),
        ("Écart moyen (min)", f"=IFERROR(AVERAGE({deviation_range}),\"—\")"),
    ]
    for i, (label, formula) in enumerate(summary_rows):
        ws.cell(row=summary_start_row + i, column=1, value=label).font = Font(name="Arial", bold=True)
        ws.cell(row=summary_start_row + i, column=2, value=formula).font = Font(name="Arial")

    for col_idx, width in enumerate([22, 20, 20, 14, 16], start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(str(output_path))

    return str(output_path)
